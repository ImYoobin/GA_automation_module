from __future__ import annotations

import unittest
from pathlib import Path
import shutil
from datetime import datetime, timedelta

from openpyxl import load_workbook

from google_ads_exporter.google_excel_builder import (
    _build_source_header_lookup,
    _build_source_header_token_lookup,
    _extract_report_day_value,
    _resolve_report_day_value,
    _resolve_missing_header_fallback,
    _resolve_source_header,
    create_unified_workbook_for_account,
)
from google_ads_exporter.models import AdsAccount, DownloadResult


class GoogleExcelBuilderHeaderMappingTests(unittest.TestCase):
    def test_resolve_source_header_matches_snake_case_tokens(self) -> None:
        source_headers = ["day", "campaign_id", "ad_group_id", "ad_name", "cost"]
        source_lookup = _build_source_header_lookup(source_headers)
        source_token_lookup = _build_source_header_token_lookup(source_headers)

        self.assertEqual(
            _resolve_source_header(
                source_lookup=source_lookup,
                source_token_lookup=source_token_lookup,
                target_header="Campaign ID",
            ),
            "campaign_id",
        )
        self.assertEqual(
            _resolve_source_header(
                source_lookup=source_lookup,
                source_token_lookup=source_token_lookup,
                target_header="Ad Group ID",
            ),
            "ad_group_id",
        )
        self.assertEqual(
            _resolve_source_header(
                source_lookup=source_lookup,
                source_token_lookup=source_token_lookup,
                target_header="Ad Name",
            ),
            "ad_name",
        )

    def test_resolve_source_header_supports_device_and_hour_aliases(self) -> None:
        source_headers = ["device_type", "hour_bucket"]
        source_lookup = _build_source_header_lookup(source_headers)
        source_token_lookup = _build_source_header_token_lookup(source_headers)

        self.assertEqual(
            _resolve_source_header(
                source_lookup=source_lookup,
                source_token_lookup=source_token_lookup,
                target_header="Device",
            ),
            "device_type",
        )
        self.assertEqual(
            _resolve_source_header(
                source_lookup=source_lookup,
                source_token_lookup=source_token_lookup,
                target_header="Hour of the Day",
            ),
            "hour_bucket",
        )

    def test_extract_report_day_value(self) -> None:
        text = 'BCG_auto_devices\n"April 13, 2026 - April 13, 2026"\nDevice\tCampaign ID'
        self.assertEqual(_extract_report_day_value(text), "2026-04-13")

        range_text = 'BCG_auto_campaignadgroup\n"March 31, 2026 - April 13, 2026"\nCampaign status'
        self.assertEqual(
            _extract_report_day_value(range_text),
            "2026-03-31 - 2026-04-13",
        )

    def test_day_fallback_allowed_only_for_specific_targets_and_single_day(self) -> None:
        self.assertEqual(
            _resolve_missing_header_fallback(
                target_key="device",
                target_header="Day",
                report_day_value="2026-04-13",
            ),
            "2026-04-13",
        )
        self.assertEqual(
            _resolve_missing_header_fallback(
                target_key="campaign_ad_group",
                target_header="Day",
                report_day_value="2026-04-13",
            ),
            "",
        )
        self.assertEqual(
            _resolve_missing_header_fallback(
                target_key="device",
                target_header="Day",
                report_day_value="2026-03-31 - 2026-04-13",
            ),
            "",
        )

    def test_resolve_report_day_value_falls_back_to_filename_for_day_targets(self) -> None:
        expected = (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")
        self.assertEqual(
            _resolve_report_day_value(
                report_day_value=None,
                target_key="hourofday",
                filename="20260415_account_FCAS_hourofday.csv",
            ),
            expected,
        )
        self.assertIsNone(
            _resolve_report_day_value(
                report_day_value=None,
                target_key="campaign_ad_group",
                filename="20260415_account_FCAS_campaign_ad_group.csv",
            )
        )

    def test_create_unified_workbook_injects_day_when_missing_in_csv(self) -> None:
        base = Path("tests/.tmp_excel_builder_mapping")
        if base.exists():
            shutil.rmtree(base, ignore_errors=True)
        try:
            csv_dir = base / "csv"
            out_dir = base / "output"
            csv_dir.mkdir(parents=True, exist_ok=True)
            out_dir.mkdir(parents=True, exist_ok=True)

            csv_path = csv_dir / "device.csv"
            csv_path.write_text(
                "\n".join(
                    [
                        "BCG_auto_devices",
                        '"April 13, 2026 - April 13, 2026"',
                        "Device,Campaign ID,Campaign,Ad group ID,Ad group,Currency code,Cost,Impr.,Trueview View,Clicks,Conversions",
                        "Mobile phones,111,Test Campaign,222,Test Ad Group,KRW,1000,10,7,3,1",
                    ]
                ),
                encoding="utf-8-sig",
            )

            account = AdsAccount(
                name="Test Account",
                cid="123-456-7890",
                cid_digits="1234567890",
            )
            results = [
                DownloadResult(
                    target_key="device",
                    success=True,
                    filename=csv_path.name,
                )
            ]
            workbook_path, _summaries = create_unified_workbook_for_account(
                account=account,
                download_results=results,
                activity_name="FCAS",
                output_dir=out_dir,
                csv_dir=csv_dir,
            )

            wb = load_workbook(workbook_path, read_only=True)
            ws = wb["devices"]
            headers = [ws.cell(row=1, column=i).value for i in range(1, 90)]
            day_col = headers.index("Day") + 1
            device_col = headers.index("Device") + 1
            campaign_id_col = headers.index("Campaign ID") + 1
            ad_group_id_col = headers.index("Ad group ID") + 1
            trueview_view_col = headers.index("trueview_views") + 1
            self.assertEqual(ws.cell(row=2, column=1).value, "2026-04-13")
            self.assertEqual(ws.cell(row=2, column=day_col).value, "2026-04-13")
            self.assertEqual(ws.cell(row=2, column=campaign_id_col).value, "111")
            self.assertEqual(ws.cell(row=2, column=ad_group_id_col).value, "222")
            self.assertEqual(ws.cell(row=2, column=device_col).value, "Mobile phones")
            self.assertEqual(str(ws.cell(row=2, column=trueview_view_col).value), "7")
        finally:
            shutil.rmtree(base, ignore_errors=True)

    def test_create_unified_workbook_emits_sheet_progress_callbacks(self) -> None:
        base = Path("tests/.tmp_excel_builder_progress")
        if base.exists():
            shutil.rmtree(base, ignore_errors=True)
        try:
            csv_dir = base / "csv"
            out_dir = base / "output"
            csv_dir.mkdir(parents=True, exist_ok=True)
            out_dir.mkdir(parents=True, exist_ok=True)

            csv_path = csv_dir / "device.csv"
            csv_path.write_text(
                "\n".join(
                    [
                        "BCG_auto_devices",
                        '"April 13, 2026 - April 13, 2026"',
                        "Device,Campaign ID,Campaign,Ad group ID,Ad group,Currency code,Cost,Impr.,Trueview View,Clicks,Conversions",
                        "Mobile phones,111,Test Campaign,222,Test Ad Group,KRW,1000,10,7,3,1",
                    ]
                ),
                encoding="utf-8-sig",
            )

            account = AdsAccount(
                name="Test Account",
                cid="123-456-7890",
                cid_digits="1234567890",
            )
            events: list[tuple[str, str, object | None]] = []

            create_unified_workbook_for_account(
                account=account,
                download_results=[
                    DownloadResult(
                        target_key="device",
                        success=True,
                        filename=csv_path.name,
                    )
                ],
                activity_name="FCAS",
                output_dir=out_dir,
                csv_dir=csv_dir,
                progress_callback=lambda target_key, stage, summary: events.append((target_key, stage, summary)),
            )

            device_events = [event for event in events if event[0] == "device"]
            self.assertEqual([event[1] for event in device_events], ["start", "completed"])
            self.assertIsNone(device_events[0][2])
            self.assertEqual(device_events[1][2].sheet_name, "devices")
        finally:
            shutil.rmtree(base, ignore_errors=True)

    def test_trueview_views_is_reported_missing_when_absent(self) -> None:
        base = Path("tests/.tmp_excel_builder_trueview_missing")
        if base.exists():
            shutil.rmtree(base, ignore_errors=True)
        try:
            csv_dir = base / "csv"
            out_dir = base / "output"
            csv_dir.mkdir(parents=True, exist_ok=True)
            out_dir.mkdir(parents=True, exist_ok=True)

            csv_path = csv_dir / "placements.csv"
            csv_path.write_text(
                "\n".join(
                    [
                        "BCG_auto_placements",
                        '"March 31, 2026 - April 13, 2026"',
                        "Day,Placement (group),Campaign,Campaign ID,Ad group,Ad group ID,Currency code,Cost,Conversions,Clicks,Viewable impr.,Impr.",
                        "2026-04-13,example.com,Campaign A,111,Group A,222,KRW,100,1,2,10,20",
                    ]
                ),
                encoding="utf-8-sig",
            )

            account = AdsAccount(
                name="Test Account",
                cid="123-456-7890",
                cid_digits="1234567890",
            )
            _workbook_path, summaries = create_unified_workbook_for_account(
                account=account,
                download_results=[
                    DownloadResult(
                        target_key="placements",
                        success=True,
                        filename=csv_path.name,
                    )
                ],
                activity_name="FCAS",
                output_dir=out_dir,
                csv_dir=csv_dir,
            )

            placements_summary = next(item for item in summaries if item.target_key == "placements")
            self.assertIn("trueview_views", placements_summary.missing_columns)
        finally:
            shutil.rmtree(base, ignore_errors=True)

    def test_campaign_ad_group_does_not_inject_day_from_date_range_line(self) -> None:
        base = Path("tests/.tmp_excel_builder_campaign_day")
        if base.exists():
            shutil.rmtree(base, ignore_errors=True)
        try:
            csv_dir = base / "csv"
            out_dir = base / "output"
            csv_dir.mkdir(parents=True, exist_ok=True)
            out_dir.mkdir(parents=True, exist_ok=True)

            csv_path = csv_dir / "campaign_ad_group.csv"
            csv_path.write_text(
                "\n".join(
                    [
                        "BCG_auto_campaignadgroup",
                        '"March 31, 2026 - April 13, 2026"',
                        "Campaign status,Campaign,Budget,Budget name,Budget type,Campaign ID,Currency code,Cost",
                        "Enabled,Test Campaign,1000,--,Daily,111,KRW,100",
                    ]
                ),
                encoding="utf-8-sig",
            )

            account = AdsAccount(
                name="Test Account",
                cid="123-456-7890",
                cid_digits="1234567890",
            )
            results = [
                DownloadResult(
                    target_key="campaign_ad_group",
                    success=True,
                    filename=csv_path.name,
                )
            ]
            workbook_path, summaries = create_unified_workbook_for_account(
                account=account,
                download_results=results,
                activity_name="FCAS",
                output_dir=out_dir,
                csv_dir=csv_dir,
            )

            wb = load_workbook(workbook_path, read_only=True)
            ws = wb["campaign_ad_group"]
            headers = [ws.cell(row=1, column=i).value for i in range(1, 90)]
            campaign_id_col = headers.index("Campaign ID") + 1
            self.assertIsNone(ws.cell(row=2, column=1).value)
            self.assertEqual(ws.cell(row=2, column=campaign_id_col).value, "111")
            self.assertIn("Day", summaries[0].missing_columns)
        finally:
            shutil.rmtree(base, ignore_errors=True)

    def test_ad_format_includes_day_and_campaign_id_with_single_day_fallback(self) -> None:
        base = Path("tests/.tmp_excel_builder_adformat")
        if base.exists():
            shutil.rmtree(base, ignore_errors=True)
        try:
            csv_dir = base / "csv"
            out_dir = base / "output"
            csv_dir.mkdir(parents=True, exist_ok=True)
            out_dir.mkdir(parents=True, exist_ok=True)

            csv_path = csv_dir / "ad_format.csv"
            csv_path.write_text(
                "\n".join(
                    [
                        "BCG_auto_adformat",
                        '"April 13, 2026 - April 13, 2026"',
                        "Ad format,Ad group,Ad group ID,Campaign ID",
                        "Skippable in-stream,Group A,222,111",
                    ]
                ),
                encoding="utf-8-sig",
            )

            account = AdsAccount(
                name="Test Account",
                cid="123-456-7890",
                cid_digits="1234567890",
            )
            results = [
                DownloadResult(
                    target_key="adformat",
                    success=True,
                    filename=csv_path.name,
                )
            ]
            workbook_path, summaries = create_unified_workbook_for_account(
                account=account,
                download_results=results,
                activity_name="FCAS",
                output_dir=out_dir,
                csv_dir=csv_dir,
            )

            wb = load_workbook(workbook_path, read_only=True)
            ws = wb["ad_format"]
            headers = [ws.cell(row=1, column=i).value for i in range(1, 120)]
            self.assertEqual(headers[:5], [
                "Day",
                "Ad format",
                "Ad group status",
                "Ad group",
                "Campaign",
            ])
            day_col = headers.index("Day") + 1
            campaign_id_col = headers.index("Campaign ID") + 1
            ad_group_id_col = headers.index("Ad group ID") + 1
            ad_group_col = headers.index("Ad group") + 1
            ad_format_col = headers.index("Ad format") + 1
            self.assertEqual(ws.cell(row=2, column=day_col).value, "2026-04-13")
            self.assertEqual(ws.cell(row=2, column=campaign_id_col).value, "111")
            self.assertEqual(ws.cell(row=2, column=ad_group_id_col).value, "222")
            self.assertEqual(ws.cell(row=2, column=ad_group_col).value, "Group A")
            self.assertEqual(ws.cell(row=2, column=ad_format_col).value, "Skippable in-stream")
            ad_format_summary = next(item for item in summaries if item.target_key == "adformat")
            self.assertNotIn("Day", ad_format_summary.missing_columns)
            self.assertNotIn("Campaign ID", ad_format_summary.missing_columns)
        finally:
            shutil.rmtree(base, ignore_errors=True)

    def test_total_row_is_dropped_from_placements_sheet(self) -> None:
        base = Path("tests/.tmp_excel_builder_total_row")
        if base.exists():
            shutil.rmtree(base, ignore_errors=True)
        try:
            csv_dir = base / "csv"
            out_dir = base / "output"
            csv_dir.mkdir(parents=True, exist_ok=True)
            out_dir.mkdir(parents=True, exist_ok=True)

            csv_path = csv_dir / "placements.csv"
            csv_path.write_text(
                "\n".join(
                    [
                        "BCG_auto_placements",
                        '"March 31, 2026 - April 13, 2026"',
                        "Day,Placement (group),Campaign,Campaign ID,Ad group,Ad group ID,Currency code,Cost,Conversions,Clicks,Viewable impr.,Impr.",
                        "2026-04-13,example.com,Campaign A,111,Group A,222,KRW,100,1,2,10,20",
                        ",Total,,,,,,100,1,2,10,20",
                    ]
                ),
                encoding="utf-8-sig",
            )

            account = AdsAccount(
                name="Test Account",
                cid="123-456-7890",
                cid_digits="1234567890",
            )
            results = [
                DownloadResult(
                    target_key="placements",
                    success=True,
                    filename=csv_path.name,
                )
            ]
            workbook_path, summaries = create_unified_workbook_for_account(
                account=account,
                download_results=results,
                activity_name="FCAS",
                output_dir=out_dir,
                csv_dir=csv_dir,
            )

            wb = load_workbook(workbook_path, read_only=True)
            ws = wb["placements"]
            self.assertEqual(ws.cell(row=2, column=2).value, "example.com")
            self.assertIsNone(ws.cell(row=3, column=2).value)
            placements_summary = next(item for item in summaries if item.target_key == "placements")
            self.assertEqual(placements_summary.written_rows, 1)
        finally:
            shutil.rmtree(base, ignore_errors=True)

    def test_placements_top100_keeps_ties_at_cutoff(self) -> None:
        base = Path("tests/.tmp_excel_builder_placements_top100")
        if base.exists():
            shutil.rmtree(base, ignore_errors=True)
        try:
            csv_dir = base / "csv"
            out_dir = base / "output"
            csv_dir.mkdir(parents=True, exist_ok=True)
            out_dir.mkdir(parents=True, exist_ok=True)

            rows: list[str] = [
                "BCG_auto_placements",
                '"March 31, 2026 - April 13, 2026"',
                "Day,Placement (group),Campaign,Campaign ID,Ad group,Ad group ID,Currency code,Cost,Conversions,Clicks,Viewable impr.,Impr.",
            ]

            # Group A: 115 rows -> keep top100 with ties at cutoff.
            # 95 rows above 100, 10 rows at 100 (cutoff tie), 10 rows at 99 (drop).
            for offset, cost in enumerate(range(200, 105, -1), start=1):
                rows.append(
                    f"2026-04-01,site-a-{offset},Campaign A,111,Group A,222,KRW,{cost},1,2,10,20"
                )
            for offset in range(1, 11):
                rows.append(
                    f"2026-04-01,site-a-cutoff-{offset},Campaign A,111,Group A,222,KRW,100,1,2,10,20"
                )
            for offset in range(1, 11):
                rows.append(
                    f"2026-04-01,site-a-drop-{offset},Campaign A,111,Group A,222,KRW,99,1,2,10,20"
                )

            # Group B: 50 rows -> keep all.
            for offset in range(1, 51):
                rows.append(
                    f"2026-04-02,site-b-{offset},Campaign B,333,Group B,444,KRW,10,1,2,10,20"
                )

            csv_path = csv_dir / "placements.csv"
            csv_path.write_text("\n".join(rows), encoding="utf-8-sig")

            account = AdsAccount(
                name="Test Account",
                cid="123-456-7890",
                cid_digits="1234567890",
            )
            results = [
                DownloadResult(
                    target_key="placements",
                    success=True,
                    filename=csv_path.name,
                )
            ]
            workbook_path, summaries = create_unified_workbook_for_account(
                account=account,
                download_results=results,
                activity_name="FCAS",
                output_dir=out_dir,
                csv_dir=csv_dir,
            )

            placements_summary = next(item for item in summaries if item.target_key == "placements")
            self.assertEqual(placements_summary.csv_rows, 165)
            self.assertEqual(placements_summary.written_rows, 155)

            wb = load_workbook(workbook_path, read_only=True)
            ws = wb["placements"]
            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            day_idx = header_row.index("Day")
            ad_group_id_idx = header_row.index("Ad group ID")
            cost_idx = header_row.index("Cost")

            kept_cost_100_group_a = 0
            kept_cost_99_group_a = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                day_value = str(row[day_idx] or "")
                ad_group_id = str(row[ad_group_id_idx] or "")
                if day_value != "2026-04-01" or ad_group_id != "222":
                    continue
                cost_text = str(row[cost_idx] or "")
                if cost_text == "100":
                    kept_cost_100_group_a += 1
                if cost_text == "99":
                    kept_cost_99_group_a += 1

            self.assertEqual(kept_cost_100_group_a, 10)
            self.assertEqual(kept_cost_99_group_a, 0)
        finally:
            shutil.rmtree(base, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
