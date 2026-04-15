from __future__ import annotations

import unittest
from unittest.mock import patch

from openpyxl import Workbook

from google_ads_exporter.google_excel_builder import (
    _build_split_sheet_name,
    _initialize_sheet_headers,
    _write_csv_rows_to_workbook,
)


class GoogleExcelBuilderSplitTests(unittest.TestCase):
    def test_write_rows_splits_into_numbered_sheets_when_limit_exceeded(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Placement"
        headers = ("Day", "Campaign ID")
        _initialize_sheet_headers(worksheet, headers)

        csv_rows = [
            {"Day": f"2026-04-{idx:02d}", "Campaign ID": str(idx)}
            for idx in range(1, 6)
        ]
        mapped = {header: header for header in headers}

        with patch("google_ads_exporter.google_excel_builder.MAX_DATA_ROWS_PER_SHEET", 2):
            written = _write_csv_rows_to_workbook(
                workbook=workbook,
                base_sheet_name="Placement",
                worksheet=worksheet,
                headers=headers,
                csv_rows=csv_rows,
                mapped_source_by_target_header=mapped,
            )

        self.assertEqual(written, 5)
        self.assertIn("Placement", workbook.sheetnames)
        self.assertIn("Placement(1)", workbook.sheetnames)
        self.assertIn("Placement(2)", workbook.sheetnames)

        self.assertEqual(workbook["Placement"].cell(row=2, column=1).value, "2026-04-01")
        self.assertEqual(workbook["Placement(1)"].cell(row=2, column=1).value, "2026-04-03")
        self.assertEqual(workbook["Placement(2)"].cell(row=2, column=1).value, "2026-04-05")

    def test_split_sheet_name_respects_excel_length_limit(self) -> None:
        name = "ABCDEFGHIJKLMNOPQRSTUVWXYZ1234567890"
        split_name = _build_split_sheet_name(name, 12)
        self.assertTrue(split_name.endswith("(12)"))
        self.assertLessEqual(len(split_name), 31)


if __name__ == "__main__":
    unittest.main()
