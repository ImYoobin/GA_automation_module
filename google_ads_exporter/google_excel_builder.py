"""Google Ads CSV -> unified Excel workbook transformer."""

from __future__ import annotations

import csv
import heapq
import io
import re
import subprocess
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

from .models import AdsAccount, DownloadResult
from .targets import TARGET_DISPLAY_NAMES, TARGET_ORDER
from .utils import sanitize_filename


@dataclass(frozen=True, slots=True)
class SheetPolicy:
    target_key: str
    sheet_name: str
    required_columns: tuple[str, ...]
    day_policy: str  # raw_only | raw_or_top_single_day


@dataclass(frozen=True, slots=True)
class CsvProcessSummary:
    target_key: str
    target_display: str
    sheet_name: str
    csv_path: str | None
    csv_rows: int
    written_rows: int
    mapped_columns: tuple[str, ...]
    missing_columns: tuple[str, ...]
    status: str
    reason: str | None = None


@dataclass(frozen=True, slots=True)
class CsvReadMetadata:
    encoding: str
    delimiter: str
    header_row_index: int
    source_headers: tuple[str, ...]
    report_day_value: str | None


@dataclass(frozen=True, slots=True)
class StreamWriteResult:
    csv_rows: int
    written_rows: int


SHEET_POLICIES: tuple[SheetPolicy, ...] = (
    SheetPolicy(
        target_key="campaign_ad_group",
        sheet_name="campaign_ad_group",
        required_columns=(
            "Day",
            "Campaign status",
            "Campaign",
            "Budget",
            "Budget name",
            "Budget type",
            "Status",
            "Status reasons",
            "Currency code",
            "Cost",
            "Impr.",
            "Unique users",
            "Avg. impr. freq. / user",
            "Clicks",
            "Conversions",
            "Results",
            "Campaign ID",
            "Bid strategy",
            "Bid strategy type",
            "Target CPA",
            "Viewable impr.",
            "Video played to 25%",
            "Video played to 50%",
            "Video played to 75%",
            "Video played to 100%",
            "Results value",
            "Video played to 25%.1",
            "Video played to 50%.1",
            "Video played to 75%.1",
            "Video played to 100%.1",
            "Viewable CTR",
            "Avg. viewable CPM",
            "Optimization score",
        ),
        day_policy="raw_only",
    ),
    SheetPolicy(
        target_key="ad",
        sheet_name="ads",
        required_columns=(
            "Day",
            "Ad status",
            "Final URL",
            "Beacon URLs",
            "Headline",
            "Long headline 1",
            "Long headline 2",
            "Long headline 3",
            "Long headline 4",
            "Long headline 5",
            "Headline 1",
            "Headline 2",
            "Headline 3",
            "Headline 4",
            "Headline 5",
            "Description 1",
            "Description 2",
            "Description 3",
            "Description 4",
            "Description 5",
            "Logo ID",
            "Business name",
            "Video",
            "Call to action text",
            "Call to action text 1",
            "Call to action text 2",
            "Call to action text 3",
            "Call to action text 4",
            "Call to action text 5",
            "Call to action headline",
            "Video ID",
            "Companion banner",
            "Ad name",
            "ad.display_url",
            "Path 1",
            "Path 2",
            "Mobile final URL",
            "Tracking template",
            "Final URL suffix",
            "Custom parameter",
            "Campaign",
            "Ad group",
            "Status",
            "Status reasons",
            "Ad type",
            "Ad strength",
            "Ad strength improvements",
            "Campaign ID",
            "Ad group ID",
            "Ad ID",
            "Currency code",
            "Cost",
            "Conversions",
            "Clicks",
            "Viewable impr.",
            "Impr.",
            "Video played to 25%",
            "Video played to 50%",
            "Video played to 75%",
            "Video played to 100%",
        ),
        day_policy="raw_only",
    ),
    SheetPolicy(
        target_key="demographics",
        sheet_name="demographics",
        required_columns=(
            "Day",
            "Age",
            "Gender",
            "Campaign",
            "Campaign ID",
            "Ad group",
            "Ad group ID",
            "Currency code",
            "Cost",
            "Conversions",
            "Clicks",
            "Impr.",
            "Video played to 25%",
            "Video played to 50%",
            "Video played to 75%",
            "Video played to 100%",
            "Avg. viewable CPM",
        ),
        day_policy="raw_only",
    ),
    SheetPolicy(
        target_key="device",
        sheet_name="devices",
        required_columns=(
            "Day",
            "Device",
            "Ad status",
            "Final URL",
            "Beacon URLs",
            "Headline",
            "Long headline 1",
            "Long headline 2",
            "Long headline 3",
            "Long headline 4",
            "Long headline 5",
            "Headline 1",
            "Headline 2",
            "Headline 3",
            "Headline 4",
            "Headline 5",
            "Description 1",
            "Description 2",
            "Description 3",
            "Description 4",
            "Description 5",
            "Logo ID",
            "Business name",
            "Video",
            "Call to action text",
            "Call to action text 1",
            "Call to action text 2",
            "Call to action text 3",
            "Call to action text 4",
            "Call to action text 5",
            "Call to action headline",
            "Video ID",
            "Companion banner",
            "Ad name",
            "ad.display_url",
            "Path 1",
            "Path 2",
            "Mobile final URL",
            "Tracking template",
            "Final URL suffix",
            "Custom parameter",
            "Campaign",
            "Ad group",
            "Status",
            "Status reasons",
            "Ad type",
            "Ad strength",
            "Ad strength improvements",
            "Campaign ID",
            "Ad group ID",
            "Ad ID",
            "Currency code",
            "Cost",
            "Conversions",
            "Clicks",
            "Viewable impr.",
            "Impr.",
            "Video played to 25%",
            "Video played to 50%",
            "Video played to 75%",
            "Video played to 100%",
        ),
        day_policy="raw_or_top_single_day",
    ),
    SheetPolicy(
        target_key="hourofday",
        sheet_name="time_of_day",
        required_columns=(
            "Day",
            "Hour of the day",
            "Ad group status",
            "Ad group",
            "Campaign",
            "Currency code",
            "Default max. CPC",
            "Max. CPV",
            "TrueView target CPV",
            "Max. CPM",
            "Target CPA",
            "Target CPM",
            "Fixed CPM",
            "Status",
            "Status reasons",
            "Campaign ID",
            "Ad group ID",
            "Currency code.1",
            "Cost",
            "Conversions",
            "Clicks",
            "Viewable impr.",
            "Impr.",
            "Video played to 25%",
            "Video played to 50%",
            "Video played to 75%",
            "Video played to 100%",
        ),
        day_policy="raw_or_top_single_day",
    ),
    SheetPolicy(
        target_key="placements",
        sheet_name="placements",
        required_columns=(
            "Day",
            "Placement (group)",
            "Campaign",
            "Campaign ID",
            "Ad group",
            "Ad group ID",
            "Currency code",
            "Cost",
            "Conversions",
            "Clicks",
            "Viewable impr.",
            "Impr.",
        ),
        day_policy="raw_only",
    ),
    SheetPolicy(
        target_key="adformat",
        sheet_name="ad_format",
        required_columns=(
            "Day",
            "Ad format",
            "Ad group status",
            "Ad group",
            "Campaign",
            "Currency code",
            "Default max. CPC",
            "Max. CPV",
            "TrueView target CPV",
            "Max. CPM",
            "Target CPA",
            "Target CPM",
            "Fixed CPM",
            "Status",
            "Status reasons",
            "Campaign ID",
            "Ad group ID",
            "Bid strategy",
            "Bid strategy type",
            "Currency code.1",
            "Cost",
            "Conversions",
            "Clicks",
            "Viewable impr.",
            "Impr.",
            "Video played to 25%",
            "Video played to 50%",
            "Video played to 75%",
            "Video played to 100%",
            "Unique users",
            "Avg. impr. freq. / user",
            "Purchase conversions",
            "All conv. value",
            "Conv. value",
        ),
        day_policy="raw_or_top_single_day",
    ),
)

POLICY_BY_TARGET: dict[str, SheetPolicy] = {policy.target_key: policy for policy in SHEET_POLICIES}
DAY_FALLBACK_ALLOWED_TARGETS = frozenset(
    policy.target_key for policy in SHEET_POLICIES if policy.day_policy == "raw_or_top_single_day"
)

HEADER_ALIAS_BY_TARGET_TOKEN: dict[str, tuple[str, ...]] = {
    "campaignname": ("campaign",),
    "adgroupname": ("adgroup",),
    "impressions": ("impr",),
    "device": ("device_type", "devicetype"),
    "devicetype": ("device",),
    "hourofday": ("hour_of_the_day", "hour", "hour_bucket", "hourofday"),
    "houroftheday": ("hour_of_day", "hour", "hour_bucket", "hourofday"),
    "adformat": ("ad_format", "adformat"),
    "placementgroup": ("placement_group", "placement"),
    "campaignstatus": ("status",),
    "adstatus": ("status",),
    "currencycode1": ("currency code",),
    "videoplayedto251": ("video played to 25%",),
    "videoplayedto501": ("video played to 50%",),
    "videoplayedto751": ("video played to 75%",),
    "videoplayedto1001": ("video played to 100%",),
    "allconvvalue": ("conv. value", "conv value"),
}

SCIENTIFIC_NOTATION_REGEX = re.compile(r"^[+-]?(?:\d+(?:\.\d+)?|\.\d+)[eE][+-]?\d+$")
SCIENTIFIC_WARNING_SAMPLE_LIMIT = 5
ILLEGAL_EXCEL_CHAR_REGEX = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
EXCEL_MAX_ROWS_PER_SHEET = 1_048_576
CSV_DATA_START_ROW = 2
MAX_DATA_ROWS_PER_SHEET = EXCEL_MAX_ROWS_PER_SHEET - CSV_DATA_START_ROW + 1
EXCEL_MAX_SHEET_NAME_LEN = 31
CSV_METADATA_PREVIEW_LINES = 120
CSV_METADATA_PREVIEW_CHARS = 262_144
TOTAL_ROW_MARKERS = {"total", "합계", "?⑷퀎"}
DATE_RANGE_LINE_REGEX = re.compile(
    r'^\s*"?([A-Za-z]+\s+\d{1,2},\s+\d{4})\s*-\s*([A-Za-z]+\s+\d{1,2},\s+\d{4})"?\s*$'
)
SINGLE_DATE_LINE_REGEX = re.compile(r'^\s*"?([A-Za-z]+\s+\d{1,2},\s+\d{4})"?\s*$')


def build_google_template_workbook() -> Workbook:
    workbook = Workbook(write_only=True)
    return workbook


def create_unified_workbook_for_account(
    *,
    account: AdsAccount,
    download_results: list[DownloadResult],
    activity_name: str = "",
    output_dir: Path,
    csv_dir: Path | None = None,
    logger=None,
) -> tuple[Path, list[CsvProcessSummary]]:
    workbook = build_google_template_workbook()
    results_by_key = {result.target_key: result for result in download_results}
    summaries: list[CsvProcessSummary] = []
    csv_source_dir = Path(csv_dir or output_dir).expanduser().resolve()

    for target_key in TARGET_ORDER:
        policy = POLICY_BY_TARGET.get(target_key)
        if policy is None:
            continue

        target_display = TARGET_DISPLAY_NAMES.get(target_key, target_key)
        result = results_by_key.get(target_key)
        summary = _write_target_to_sheet(
            workbook=workbook,
            policy=policy,
            target_display=target_display,
            result=result,
            csv_source_dir=csv_source_dir,
            logger=logger,
        )
        summaries.append(summary)

    run_date = datetime.now().strftime("%Y%m%d")
    activity_fragment = str(activity_name or "").strip()
    if activity_fragment:
        output_stem = f"{run_date}_{account.name}_{account.cid_digits}_{activity_fragment}_Google_Unified.xlsx"
    else:
        output_stem = f"{run_date}_{account.name}_{account.cid_digits}_Google_Unified.xlsx"
    output_name = sanitize_filename(output_stem)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / output_name
    temp_output_path = output_dir / f".{output_name}.tmp"
    if temp_output_path.exists():
        temp_output_path.unlink()
    workbook.save(temp_output_path)
    temp_output_path.replace(output_path)
    return output_path, summaries


def open_file_in_explorer(file_path: Path, logger=None) -> bool:
    target = Path(file_path).resolve()
    try:
        subprocess.Popen(["explorer", f"/select,{target}"])  # noqa: S603
        return True
    except Exception as exc:  # noqa: BLE001
        if logger:
            logger.warning("explorer select failed path=%s reason=%s", target, exc)

    try:
        subprocess.Popen(["explorer", str(target.parent)])  # noqa: S603
        return True
    except Exception as exc:  # noqa: BLE001
        if logger:
            logger.warning("explorer open folder failed path=%s reason=%s", target.parent, exc)
    return False


def summaries_as_rows(
    summaries: list[CsvProcessSummary],
    account_label: str,
    activity_name: str = "",
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for item in summaries:
        rows.append(
            {
                "account": account_label,
                "activity": activity_name,
                "target_key": item.target_key,
                "target_display": item.target_display,
                "sheet_name": item.sheet_name,
                "csv_path": item.csv_path or "",
                "csv_rows": item.csv_rows,
                "written_rows": item.written_rows,
                "mapped_columns": len(item.mapped_columns),
                "missing_columns": len(item.missing_columns),
                "status": item.status,
                "reason": item.reason or "",
            }
        )
    return rows


def _write_target_to_sheet(
    *,
    workbook: Workbook,
    policy: SheetPolicy,
    target_display: str,
    result: DownloadResult | None,
    csv_source_dir: Path,
    logger=None,
) -> CsvProcessSummary:
    required_headers = policy.required_columns

    if not result:
        return CsvProcessSummary(
            target_key=policy.target_key,
            target_display=target_display,
            sheet_name=policy.sheet_name,
            csv_path=None,
            csv_rows=0,
            written_rows=0,
            mapped_columns=tuple(),
            missing_columns=required_headers,
            status="failed",
            reason="download result missing",
        )

    if not result.success or not result.filename:
        return CsvProcessSummary(
            target_key=policy.target_key,
            target_display=target_display,
            sheet_name=policy.sheet_name,
            csv_path=None,
            csv_rows=0,
            written_rows=0,
            mapped_columns=tuple(),
            missing_columns=required_headers,
            status="failed",
            reason=result.reason or "download failed",
        )

    csv_path = (csv_source_dir / result.filename).resolve()
    if not csv_path.exists():
        return CsvProcessSummary(
            target_key=policy.target_key,
            target_display=target_display,
            sheet_name=policy.sheet_name,
            csv_path=str(csv_path),
            csv_rows=0,
            written_rows=0,
            mapped_columns=tuple(),
            missing_columns=required_headers,
            status="failed",
            reason="csv file not found",
        )

    try:
        metadata = _read_csv_metadata(csv_path)
    except Exception as exc:  # noqa: BLE001
        if logger:
            logger.warning("csv read failed path=%s reason=%s", csv_path, exc)
        return CsvProcessSummary(
            target_key=policy.target_key,
            target_display=target_display,
            sheet_name=policy.sheet_name,
            csv_path=str(csv_path),
            csv_rows=0,
            written_rows=0,
            mapped_columns=tuple(),
            missing_columns=required_headers,
            status="failed",
            reason=f"csv read failed: {exc}",
        )

    source_headers = list(metadata.source_headers)
    report_day_value = _resolve_report_day_value(
        report_day_value=metadata.report_day_value,
        target_key=policy.target_key,
        filename=result.filename,
        logger=logger,
    )

    source_lookup = _build_source_header_lookup(source_headers)
    source_token_lookup = _build_source_header_token_lookup(source_headers)
    ordered_source = _build_ordered_source_columns(source_headers)

    mapped_columns: list[str] = []
    missing_columns: list[str] = []
    mapped_source_by_target_header: dict[str, str] = {}
    default_value_by_target_header: dict[str, str] = {}
    consumed_source_headers: set[str] = set()

    for required_header in required_headers:
        source_header = _resolve_source_header(
            source_lookup=source_lookup,
            source_token_lookup=source_token_lookup,
            target_header=required_header,
        )
        if source_header:
            mapped_columns.append(required_header)
            mapped_source_by_target_header[required_header] = source_header
            consumed_source_headers.add(_normalize_header(source_header))
            continue

        mapped_source_by_target_header[required_header] = ""
        fallback = _resolve_missing_header_fallback(
            target_key=policy.target_key,
            target_header=required_header,
            report_day_value=report_day_value,
        )
        if fallback:
            mapped_columns.append(required_header)
            default_value_by_target_header[required_header] = fallback
        else:
            missing_columns.append(required_header)

    final_headers = list(required_headers)
    required_header_set = set(required_headers)
    for canonical_header, source_header in ordered_source:
        if canonical_header in required_header_set:
            continue
        if canonical_header in consumed_source_headers:
            continue
        if canonical_header in mapped_source_by_target_header:
            continue
        final_headers.append(canonical_header)
        mapped_source_by_target_header[canonical_header] = source_header

    worksheet = workbook.create_sheet(policy.sheet_name)
    _initialize_sheet_headers(worksheet, tuple(final_headers))

    stream_result = _write_csv_stream_to_workbook(
        workbook=workbook,
        base_sheet_name=policy.sheet_name,
        worksheet=worksheet,
        target_key=policy.target_key,
        csv_path=csv_path,
        csv_meta=metadata,
        headers=tuple(final_headers),
        source_headers=tuple(source_headers),
        mapped_source_by_target_header=mapped_source_by_target_header,
        default_value_by_target_header=default_value_by_target_header,
        report_day_value=report_day_value,
        logger=logger,
    )

    if logger and missing_columns:
        logger.warning(
            "required header missing in csv | target=%s | sheet=%s | csv=%s | missing=%s",
            policy.target_key,
            policy.sheet_name,
            csv_path,
            missing_columns,
        )

    return CsvProcessSummary(
        target_key=policy.target_key,
        target_display=target_display,
        sheet_name=policy.sheet_name,
        csv_path=str(csv_path),
        csv_rows=stream_result.csv_rows,
        written_rows=stream_result.written_rows,
        mapped_columns=tuple(mapped_columns),
        missing_columns=tuple(missing_columns),
        status="excel_written",
    )


def _read_csv_metadata(path: Path) -> CsvReadMetadata:
    encoding = _detect_csv_encoding(path)
    preview_text = _read_csv_preview(path=path, encoding=encoding)
    candidates = _build_delimiter_candidates(preview_text)
    if not candidates:
        candidates = (",",)

    best_delimiter = candidates[0]
    best_header_row_index = 0
    best_headers: list[str] = []
    best_score = -1
    for delimiter in candidates:
        header_row_index = _detect_header_row_index(text=preview_text, delimiter=delimiter, scan_limit=60)
        headers = _parse_header_row_from_preview(
            text=preview_text,
            delimiter=delimiter,
            header_row_index=header_row_index,
        )
        score = _score_parsed_csv(headers=headers, rows=[])
        if score > best_score:
            best_score = score
            best_delimiter = delimiter
            best_header_row_index = header_row_index
            best_headers = headers

    if not best_headers:
        raise RuntimeError("header row not found from CSV preview")

    report_day_value = _extract_report_day_value(preview_text)
    return CsvReadMetadata(
        encoding=encoding,
        delimiter=best_delimiter,
        header_row_index=best_header_row_index,
        source_headers=tuple(best_headers),
        report_day_value=report_day_value,
    )


def _read_csv_preview(*, path: Path, encoding: str) -> str:
    lines: list[str] = []
    total_chars = 0
    with path.open("r", encoding=encoding, errors="replace", newline="") as fp:
        for _ in range(CSV_METADATA_PREVIEW_LINES):
            line = fp.readline()
            if not line:
                break
            lines.append(line)
            total_chars += len(line)
            if total_chars >= CSV_METADATA_PREVIEW_CHARS:
                break
    return "".join(lines)


def _detect_csv_encoding(path: Path) -> str:
    with path.open("rb") as fp:
        raw = fp.read(4096)
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if raw.startswith(b"\xff\xfe"):
        return "utf-16"
    if raw.startswith(b"\xfe\xff"):
        return "utf-16"
    if b"\x00" in raw:
        return "utf-16"

    for encoding in ("utf-8-sig", "utf-8", "cp949"):
        try:
            raw.decode(encoding)
            return encoding
        except Exception:  # noqa: BLE001
            continue
    return "utf-8-sig"


def _parse_header_row_from_preview(*, text: str, delimiter: str, header_row_index: int) -> list[str]:
    with io.StringIO(text) as fp:
        reader = csv.reader(fp, delimiter=delimiter)
        for _ in range(max(0, header_row_index)):
            next(reader, None)
        header_row = next(reader, [])
    return [str(name or "").strip() for name in header_row]


def _parse_csv_text(text: str, *, delimiter: str) -> tuple[list[dict[str, str]], list[str]]:
    header_row_index = _detect_header_row_index(text=text, delimiter=delimiter, scan_limit=60)
    rows: list[dict[str, str]] = []
    with io.StringIO(text) as fp:
        reader = csv.reader(fp, delimiter=delimiter)

        for _ in range(header_row_index):
            next(reader, None)

        header_row = next(reader, [])
        headers = [str(name or "").strip() for name in header_row]

        for raw_row in reader:
            if not raw_row:
                continue
            normalized = _build_row_dict(headers=headers, raw_row=raw_row)
            if _is_row_completely_blank(normalized):
                continue
            if _is_total_row(normalized):
                continue
            rows.append(normalized)
    return rows, headers


def _build_row_dict(*, headers: list[str], raw_row: list[str]) -> dict[str, str]:
    normalized: dict[str, str] = {}
    for idx, header in enumerate(headers):
        key = str(header or "").strip()
        value = str(raw_row[idx] if idx < len(raw_row) else "").strip()
        normalized[key] = value
    return normalized


def _detect_header_row_index(*, text: str, delimiter: str, scan_limit: int) -> int:
    preview_rows: list[list[str]] = []
    with io.StringIO(text) as fp:
        reader = csv.reader(fp, delimiter=delimiter)
        for row_idx, row in enumerate(reader):
            preview_rows.append([str(cell or "").strip() for cell in row])
            if row_idx + 1 >= scan_limit:
                break

    best_index = 0
    best_score = -1
    for row_idx, row in enumerate(preview_rows):
        score = _score_header_candidate(row)
        if score > best_score:
            best_index = row_idx
            best_score = score
    return best_index


def _score_header_candidate(row: list[str]) -> int:
    non_empty_cells = [cell for cell in row if cell]
    if len(non_empty_cells) < 2:
        return -1

    tokenized = {_header_token(cell) for cell in non_empty_cells}
    anchors = (
        "day",
        "campaign",
        "campaignid",
        "adgroup",
        "adgroupid",
        "adid",
        "cost",
        "impr",
        "clicks",
        "conversions",
        "currencycode",
        "placementgroup",
        "houroftheday",
        "adformat",
    )
    anchor_hits = sum(1 for anchor in anchors if anchor in tokenized)
    return (len(non_empty_cells) * 10) + (anchor_hits * 25)


def _header_token(value: str) -> str:
    return "".join(char for char in str(value or "").strip().lower() if char.isalnum())


def _build_delimiter_candidates(text: str) -> tuple[str, ...]:
    candidates: list[str] = []
    sample = _preview_csv_text(text, max_lines=60, max_chars=32768)
    sniffed = _sniff_delimiter(sample)
    if sniffed:
        candidates.append(sniffed)
    for delimiter in ("\t", ",", ";", "|"):
        if delimiter not in candidates:
            candidates.append(delimiter)
    return tuple(candidates)


def _sniff_delimiter(sample: str) -> str | None:
    if not sample.strip():
        return None
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="\t,;|")
        return str(dialect.delimiter)
    except Exception:  # noqa: BLE001
        return _heuristic_delimiter(sample)


def _heuristic_delimiter(sample: str) -> str | None:
    lines = [line for line in sample.splitlines() if line.strip()]
    if not lines:
        return None
    best_delimiter: str | None = None
    best_score = -1
    for delimiter in ("\t", ",", ";", "|"):
        counts = [line.count(delimiter) for line in lines[:20]]
        non_zero = [count for count in counts if count > 0]
        if not non_zero:
            continue
        score = (sum(non_zero) * 10) - (max(non_zero) - min(non_zero))
        if score > best_score:
            best_score = score
            best_delimiter = delimiter
    return best_delimiter


def _preview_csv_text(text: str, *, max_lines: int, max_chars: int) -> str:
    if len(text) <= max_chars:
        return "\n".join(text.splitlines()[:max_lines])

    preview = text[:max_chars]
    lines = preview.splitlines()
    return "\n".join(lines[:max_lines])


def _score_parsed_csv(*, headers: list[str], rows: list[dict[str, str]]) -> int:
    header_count = sum(1 for header in headers if str(header or "").strip())
    if header_count == 0:
        return -1
    row_count = len(rows)
    return (header_count * 100000) + min(row_count, 99999)


def _write_csv_stream_to_workbook(
    *,
    workbook: Workbook,
    base_sheet_name: str,
    worksheet,
    target_key: str,
    csv_path: Path,
    csv_meta: CsvReadMetadata,
    headers: tuple[str, ...],
    source_headers: tuple[str, ...],
    mapped_source_by_target_header: dict[str, str],
    default_value_by_target_header: dict[str, str] | None = None,
    report_day_value: str | None = None,
    logger=None,
) -> StreamWriteResult:
    default_lookup = default_value_by_target_header or {}
    source_index_by_header: dict[str, int] = {}
    for idx, header in enumerate(source_headers):
        key = str(header or "").strip()
        if key and key not in source_index_by_header:
            source_index_by_header[key] = idx

    mapped_source_index_by_target_header: dict[str, int] = {}
    for target_header, source_header in mapped_source_by_target_header.items():
        source_index = source_index_by_header.get(source_header or "")
        if source_index is not None:
            mapped_source_index_by_target_header[target_header] = source_index

    column_plan: list[tuple[int, str, bool]] = []
    for header in headers:
        source_index = mapped_source_index_by_target_header.get(header, -1)
        default_value = _as_trimmed_string(default_lookup.get(header, ""))
        is_day_header = _header_token(header) == "day"
        column_plan.append((source_index, default_value, is_day_header))
    single_day_fill = str(report_day_value or "").strip() if _is_single_day_report_value(report_day_value) else ""

    if target_key == "placements":
        return _write_placements_top_rows_stream(
            workbook=workbook,
            base_sheet_name=base_sheet_name,
            worksheet=worksheet,
            csv_path=csv_path,
            csv_meta=csv_meta,
            headers=headers,
            source_headers=source_headers,
            mapped_source_index_by_target_header=mapped_source_index_by_target_header,
            column_plan=column_plan,
            logger=logger,
        )

    scientific_headers = [
        (idx, header)
        for idx, header in enumerate(source_headers)
        if _is_id_like_header(header)
    ]
    track_scientific = (
        bool(logger)
        and bool(scientific_headers)
        and csv_path.stat().st_size <= 80 * 1024 * 1024
    )
    scientific_count = 0
    scientific_samples: list[str] = []

    csv_rows = 0
    written_rows = 0
    part = 0
    rows_in_sheet = 0
    target_ws = worksheet
    cell_mode = callable(getattr(target_ws, "cell", None))

    with csv_path.open("r", encoding=csv_meta.encoding, errors="replace", newline="") as fp:
        reader = csv.reader(fp, delimiter=csv_meta.delimiter)

        for _ in range(max(0, csv_meta.header_row_index)):
            next(reader, None)
        next(reader, None)  # header row

        for raw_row in reader:
            if _is_raw_row_blank(raw_row):
                continue
            if _is_raw_row_total(raw_row):
                continue
            csv_rows += 1

            if rows_in_sheet >= MAX_DATA_ROWS_PER_SHEET:
                part += 1
                rows_in_sheet = 0
                target_sheet_name = _build_split_sheet_name(base_sheet_name, part)
                target_ws = workbook.create_sheet(target_sheet_name)
                _initialize_sheet_headers(target_ws, headers)
                cell_mode = callable(getattr(target_ws, "cell", None))

            row_values: list[str] = []
            for source_index, default_value, is_day_header in column_plan:
                if source_index < 0:
                    value = default_value
                else:
                    raw_value = raw_row[source_index] if source_index < len(raw_row) else ""
                    value = _as_trimmed_string(raw_value)
                    if not value:
                        value = default_value
                    if not value and is_day_header and single_day_fill:
                        value = single_day_fill
                row_values.append(value)

            if cell_mode:
                row_idx = CSV_DATA_START_ROW + rows_in_sheet
                for col_idx, value in enumerate(row_values, start=1):
                    target_ws.cell(row=row_idx, column=col_idx, value=value)
            else:
                target_ws.append(row_values)

            if track_scientific:
                for source_index, header in scientific_headers:
                    raw_value = raw_row[source_index] if source_index < len(raw_row) else ""
                    text_value = _as_trimmed_string(raw_value)
                    if not text_value:
                        continue
                    if SCIENTIFIC_NOTATION_REGEX.match(text_value):
                        scientific_count += 1
                        if len(scientific_samples) < SCIENTIFIC_WARNING_SAMPLE_LIMIT:
                            scientific_samples.append(f"r{csv_rows + 1}:{header}={text_value}")

            rows_in_sheet += 1
            written_rows += 1

    if logger and part > 0:
        logger.info(
            "sheet split written | base_sheet=%s | total_rows=%s | parts=%s",
            base_sheet_name,
            written_rows,
            part + 1,
        )
    if logger and scientific_count > 0:
        logger.warning(
            "scientific notation detected in ID-like fields | csv=%s | count=%s | sample=%s",
            csv_path,
            scientific_count,
            ", ".join(scientific_samples),
        )

    return StreamWriteResult(csv_rows=csv_rows, written_rows=written_rows)


def _write_placements_top_rows_stream(
    *,
    workbook: Workbook,
    base_sheet_name: str,
    worksheet,
    csv_path: Path,
    csv_meta: CsvReadMetadata,
    headers: tuple[str, ...],
    source_headers: tuple[str, ...],
    mapped_source_index_by_target_header: dict[str, int],
    column_plan: list[tuple[int, str, bool]],
    logger=None,
) -> StreamWriteResult:
    day_idx = _resolve_source_index_for_header(
        source_headers=source_headers,
        mapped_source_index_by_target_header=mapped_source_index_by_target_header,
        target_header="Day",
    )
    ad_group_id_idx = _resolve_source_index_for_header(
        source_headers=source_headers,
        mapped_source_index_by_target_header=mapped_source_index_by_target_header,
        target_header="Ad group ID",
    )
    cost_idx = _resolve_source_index_for_header(
        source_headers=source_headers,
        mapped_source_index_by_target_header=mapped_source_index_by_target_header,
        target_header="Cost",
    )

    group_counts: dict[tuple[str, str], int] = defaultdict(int)
    group_top_costs: dict[tuple[str, str], list[float]] = defaultdict(list)
    csv_rows = 0

    for raw_row in _iter_csv_raw_rows(path=csv_path, meta=csv_meta):
        csv_rows += 1
        group_key = _placement_group_key(raw_row=raw_row, day_idx=day_idx, ad_group_id_idx=ad_group_id_idx)
        group_counts[group_key] += 1
        parsed_cost = _parse_cost_value(_value_by_index(raw_row, cost_idx))
        cost_value = float("-inf") if parsed_cost is None else parsed_cost
        heap = group_top_costs[group_key]
        heapq.heappush(heap, cost_value)
        if len(heap) > 100:
            heapq.heappop(heap)

    group_cutoff_by_key: dict[tuple[str, str], float | None] = {}
    for group_key, count in group_counts.items():
        if count <= 100:
            group_cutoff_by_key[group_key] = None
            continue
        heap = group_top_costs.get(group_key, [])
        group_cutoff_by_key[group_key] = heap[0] if heap else float("-inf")

    written_rows = 0
    part = 0
    rows_in_sheet = 0
    target_ws = worksheet
    cell_mode = callable(getattr(target_ws, "cell", None))

    for raw_row in _iter_csv_raw_rows(path=csv_path, meta=csv_meta):
        group_key = _placement_group_key(raw_row=raw_row, day_idx=day_idx, ad_group_id_idx=ad_group_id_idx)
        cutoff = group_cutoff_by_key.get(group_key)
        if cutoff is not None:
            parsed_cost = _parse_cost_value(_value_by_index(raw_row, cost_idx))
            cost_value = float("-inf") if parsed_cost is None else parsed_cost
            if cost_value < cutoff:
                continue

        if rows_in_sheet >= MAX_DATA_ROWS_PER_SHEET:
            part += 1
            rows_in_sheet = 0
            target_sheet_name = _build_split_sheet_name(base_sheet_name, part)
            target_ws = workbook.create_sheet(target_sheet_name)
            _initialize_sheet_headers(target_ws, headers)
            cell_mode = callable(getattr(target_ws, "cell", None))

        row_values: list[str] = []
        for source_index, default_value, _is_day_header in column_plan:
            if source_index < 0:
                value = default_value
            else:
                raw_value = _value_by_index(raw_row, source_index)
                value = _as_trimmed_string(raw_value)
                if not value:
                    value = default_value
            row_values.append(value)

        if cell_mode:
            row_idx = CSV_DATA_START_ROW + rows_in_sheet
            for col_idx, value in enumerate(row_values, start=1):
                target_ws.cell(row=row_idx, column=col_idx, value=value)
        else:
            target_ws.append(row_values)

        rows_in_sheet += 1
        written_rows += 1

    if logger:
        group_total = len(group_counts)
        dropped_rows = max(0, csv_rows - written_rows)
        logger.info(
            "placements top100 filter applied | groups=%s | input_rows=%s | written_rows=%s | dropped_rows=%s",
            group_total,
            csv_rows,
            written_rows,
            dropped_rows,
        )
        if part > 0:
            logger.info(
                "sheet split written | base_sheet=%s | total_rows=%s | parts=%s",
                base_sheet_name,
                written_rows,
                part + 1,
            )

    return StreamWriteResult(csv_rows=csv_rows, written_rows=written_rows)


def _iter_csv_raw_rows(*, path: Path, meta: CsvReadMetadata):
    with path.open("r", encoding=meta.encoding, errors="replace", newline="") as fp:
        reader = csv.reader(fp, delimiter=meta.delimiter)
        for _ in range(max(0, meta.header_row_index)):
            next(reader, None)
        next(reader, None)  # header
        for raw_row in reader:
            if _is_raw_row_blank(raw_row):
                continue
            if _is_raw_row_total(raw_row):
                continue
            yield raw_row


def _resolve_source_index_for_header(
    *,
    source_headers: tuple[str, ...],
    mapped_source_index_by_target_header: dict[str, int],
    target_header: str,
) -> int:
    direct = mapped_source_index_by_target_header.get(target_header)
    if direct is not None:
        return direct

    lookup = _build_source_header_lookup(list(source_headers))
    token_lookup = _build_source_header_token_lookup(list(source_headers))
    resolved = _resolve_source_header(
        source_lookup=lookup,
        source_token_lookup=token_lookup,
        target_header=target_header,
    )
    if not resolved:
        return -1
    key = str(resolved or "").strip()
    for idx, header in enumerate(source_headers):
        if str(header or "").strip() == key:
            return idx
    return -1


def _placement_group_key(*, raw_row: list[str], day_idx: int, ad_group_id_idx: int) -> tuple[str, str]:
    day_value = _as_trimmed_string(_value_by_index(raw_row, day_idx))
    ad_group_id_value = _as_trimmed_string(_value_by_index(raw_row, ad_group_id_idx))
    return day_value, ad_group_id_value


def _value_by_index(raw_row: list[str], index: int) -> str:
    if index < 0 or index >= len(raw_row):
        return ""
    return str(raw_row[index] or "")


def _parse_cost_value(raw_value: str) -> float | None:
    text = _as_trimmed_string(raw_value)
    if not text:
        return None

    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative = True
        text = text[1:-1].strip()

    cleaned = re.sub(r"[^0-9.\-]", "", text.replace(",", ""))
    if not cleaned:
        return None
    try:
        value = float(cleaned)
    except Exception:  # noqa: BLE001
        return None
    if negative:
        value = -abs(value)
    return value


def _write_csv_rows_to_sheet(
    *,
    worksheet,
    headers: tuple[str, ...],
    csv_rows: list[dict[str, str]],
    mapped_source_by_target_header: dict[str, str],
    default_value_by_target_header: dict[str, str] | None = None,
    start_row: int = CSV_DATA_START_ROW,
    start_index: int = 0,
    end_index: int | None = None,
) -> int:
    default_lookup = default_value_by_target_header or {}
    stop = len(csv_rows) if end_index is None else max(0, min(end_index, len(csv_rows)))
    begin = max(0, min(start_index, stop))
    written_rows = 0
    for offset, csv_row in enumerate(csv_rows[begin:stop], start=0):
        row_idx = start_row + offset
        for col_idx, header in enumerate(headers, start=1):
            source_header = mapped_source_by_target_header.get(header, "")
            if source_header:
                value = _as_trimmed_string(csv_row.get(source_header, ""))
            else:
                value = _as_trimmed_string(default_lookup.get(header, ""))
            worksheet.cell(row=row_idx, column=col_idx, value=value)
        written_rows += 1
    return written_rows


def _write_csv_rows_to_workbook(
    *,
    workbook: Workbook,
    base_sheet_name: str,
    worksheet,
    headers: tuple[str, ...],
    csv_rows: list[dict[str, str]],
    mapped_source_by_target_header: dict[str, str],
    default_value_by_target_header: dict[str, str] | None = None,
    logger=None,
) -> int:
    if not csv_rows:
        return 0

    total = len(csv_rows)
    if total <= MAX_DATA_ROWS_PER_SHEET:
        return _write_csv_rows_to_sheet(
            worksheet=worksheet,
            headers=headers,
            csv_rows=csv_rows,
            mapped_source_by_target_header=mapped_source_by_target_header,
            default_value_by_target_header=default_value_by_target_header,
        )

    written_rows = 0
    start = 0
    part = 0
    while start < total:
        end = min(start + MAX_DATA_ROWS_PER_SHEET, total)
        if part == 0:
            target_ws = worksheet
            target_sheet_name = base_sheet_name
        else:
            target_sheet_name = _build_split_sheet_name(base_sheet_name, part)
            target_ws = workbook.create_sheet(target_sheet_name)
            _initialize_sheet_headers(target_ws, headers)

        written_rows += _write_csv_rows_to_sheet(
            worksheet=target_ws,
            headers=headers,
            csv_rows=csv_rows,
            mapped_source_by_target_header=mapped_source_by_target_header,
            default_value_by_target_header=default_value_by_target_header,
            start_row=CSV_DATA_START_ROW,
            start_index=start,
            end_index=end,
        )
        start = end
        part += 1

    if logger:
        logger.info(
            "sheet split written | base_sheet=%s | total_rows=%s | parts=%s",
            base_sheet_name,
            total,
            part,
        )
    return written_rows


def _initialize_sheet_headers(worksheet, headers: tuple[str, ...]) -> None:
    if callable(getattr(worksheet, "cell", None)):
        for col_idx, header in enumerate(headers, start=1):
            worksheet.cell(row=1, column=col_idx, value=header)
            worksheet.cell(row=2, column=col_idx, value="")
        return
    worksheet.append(list(headers))


def _build_split_sheet_name(base_sheet_name: str, part_index: int) -> str:
    suffix = f"({part_index})"
    max_base_len = EXCEL_MAX_SHEET_NAME_LEN - len(suffix)
    trimmed_base = str(base_sheet_name or "").strip()[:max_base_len]
    if not trimmed_base:
        trimmed_base = "Sheet"
    return f"{trimmed_base}{suffix}"


def _as_trimmed_string(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        text = value.strip()
    else:
        text = str(value).strip()
    if not text:
        return ""
    if ILLEGAL_EXCEL_CHAR_REGEX.search(text) is None:
        return text
    return ILLEGAL_EXCEL_CHAR_REGEX.sub("", text)


def _is_row_completely_blank(row: dict[str, str]) -> bool:
    for value in row.values():
        if str(value or "").strip():
            return False
    return True


def _is_total_row(row: dict[str, str]) -> bool:
    for value in row.values():
        text = str(value or "").strip().lower()
        if text in TOTAL_ROW_MARKERS:
            return True
    return False


def _is_raw_row_blank(raw_row: list[str]) -> bool:
    for value in raw_row:
        if str(value or "").strip():
            return False
    return True


def _is_raw_row_total(raw_row: list[str]) -> bool:
    for value in raw_row:
        text = str(value or "").strip().lower()
        if text in TOTAL_ROW_MARKERS:
            return True
    return False


def _warn_scientific_notation_values(
    *,
    csv_path: Path,
    source_headers: list[str],
    csv_rows: list[dict[str, str]],
    logger=None,
) -> None:
    if logger is None:
        return

    id_headers = [header for header in source_headers if _is_id_like_header(header)]
    if not id_headers:
        return

    detections: list[tuple[int, str, str]] = []
    for row_idx, row in enumerate(csv_rows, start=2):
        for header in id_headers:
            raw = _as_trimmed_string(row.get(header, ""))
            if not raw:
                continue
            if SCIENTIFIC_NOTATION_REGEX.match(raw):
                detections.append((row_idx, header, raw))

    if not detections:
        return

    sample = ", ".join(
        [
            f"r{row_idx}:{header}={value}"
            for row_idx, header, value in detections[:SCIENTIFIC_WARNING_SAMPLE_LIMIT]
        ]
    )
    logger.warning(
        "scientific notation detected in ID-like fields | csv=%s | count=%s | sample=%s",
        csv_path,
        len(detections),
        sample,
    )


def _is_id_like_header(header: str) -> bool:
    normalized = _normalize_header(header)
    return normalized.endswith("id")


def _build_source_header_lookup(headers: list[str]) -> dict[str, str]:
    lookup: dict[str, str] = {}
    for header in headers:
        key = _normalize_header(header)
        if key and key not in lookup:
            lookup[key] = header
    return lookup


def _build_source_header_token_lookup(headers: list[str]) -> dict[str, str]:
    lookup: dict[str, str] = {}
    for header in headers:
        token = _header_token(header)
        if token and token not in lookup:
            lookup[token] = header
    return lookup


def _build_ordered_source_columns(headers: list[str]) -> list[tuple[str, str]]:
    ordered: list[tuple[str, str]] = []
    seen: set[str] = set()
    for header in headers:
        canonical = _normalize_header(header)
        if not canonical:
            continue
        if canonical in seen:
            continue
        seen.add(canonical)
        ordered.append((canonical, header))
    return ordered


def _resolve_source_header(
    *,
    source_lookup: dict[str, str],
    source_token_lookup: dict[str, str],
    target_header: str,
) -> str:
    exact = source_lookup.get(_normalize_header(target_header), "")
    if exact:
        return exact

    target_token = _header_token(target_header)
    token_match = source_token_lookup.get(target_token, "")
    if token_match:
        return token_match

    aliases = HEADER_ALIAS_BY_TARGET_TOKEN.get(target_token, tuple())
    for alias in aliases:
        alias_exact = source_lookup.get(_normalize_header(alias), "")
        if alias_exact:
            return alias_exact
        alias_match = source_token_lookup.get(_header_token(alias), "")
        if alias_match:
            return alias_match
    return ""


def _resolve_missing_header_fallback(
    *,
    target_key: str,
    target_header: str,
    report_day_value: str | None,
) -> str:
    if _header_token(target_header) != "day":
        return ""
    if target_key not in DAY_FALLBACK_ALLOWED_TARGETS:
        return ""
    if not _is_single_day_report_value(report_day_value):
        return ""
    if report_day_value:
        return report_day_value
    return ""


def _resolve_report_day_value(
    *,
    report_day_value: str | None,
    target_key: str,
    filename: str | None,
    logger=None,
) -> str | None:
    raw_value = str(report_day_value or "").strip()
    if raw_value:
        return raw_value
    if target_key not in DAY_FALLBACK_ALLOWED_TARGETS:
        return None

    fallback = _infer_yesterday_from_execution_date()
    if fallback and logger:
        logger.info(
            "day fallback inferred from execution date | target=%s | filename=%s | day=%s",
            target_key,
            filename,
            fallback,
        )
    return fallback


def _infer_yesterday_from_execution_date() -> str:
    run_day = datetime.now()
    inferred = run_day - timedelta(days=1)
    return inferred.strftime("%Y-%m-%d")


def _extract_report_day_value(text: str) -> str | None:
    non_empty_lines = [line.strip() for line in text.splitlines() if str(line or "").strip()]
    if not non_empty_lines:
        return None

    scan_window = non_empty_lines[:10]
    for raw_line in scan_window:
        parsed = _parse_report_date_line(raw_line)
        if parsed:
            return parsed
    return None


def _parse_report_date_line(raw_line: str) -> str | None:
    cleaned = str(raw_line or "").strip()
    if not cleaned:
        return None

    range_match = DATE_RANGE_LINE_REGEX.match(cleaned)
    if range_match:
        start = _parse_month_day_year(range_match.group(1))
        end = _parse_month_day_year(range_match.group(2))
        if not start or not end:
            return None
        if start == end:
            return start.strftime("%Y-%m-%d")
        return f"{start.strftime('%Y-%m-%d')} - {end.strftime('%Y-%m-%d')}"

    single_match = SINGLE_DATE_LINE_REGEX.match(cleaned)
    if single_match:
        only = _parse_month_day_year(single_match.group(1))
        if only:
            return only.strftime("%Y-%m-%d")
    return None


def _parse_month_day_year(raw_value: str):
    text = str(raw_value or "").strip()
    if not text:
        return None
    try:
        return datetime.strptime(text, "%B %d, %Y")
    except Exception:  # noqa: BLE001
        return None


def _is_single_day_report_value(value: str | None) -> bool:
    if not value:
        return False
    return " - " not in str(value)


def _normalize_header(value: str) -> str:
    raw = str(value or "").strip().lower()
    if not raw:
        return ""
    tokens = re.findall(r"[a-z0-9]+", raw)
    return "_".join(tokens)
